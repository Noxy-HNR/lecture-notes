"""Startup must retain the menu, share imports, and report failures before capture."""
import json
import subprocess
import sys
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest

import main
import transcribe

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def isolated_preparation(monkeypatch):
    monkeypatch.setattr(transcribe, "_dependency_future", None)
    monkeypatch.setattr(transcribe, "_model", None)


def test_menu_prepares_imports_before_waiting_without_loading_weights(monkeypatch, isolated_preparation):
    entered = threading.Event()
    release = threading.Event()
    dependencies = (object(), object(), object())

    def imports():
        entered.set()
        assert release.wait(3)
        return dependencies

    def answer(prompt):
        assert entered.wait(3)
        assert transcribe._model is None
        return "1"

    monkeypatch.setattr(transcribe, "_import_dependencies", imports)
    monkeypatch.setattr("builtins.input", answer)
    try:
        assert main.choose_action(None) == "record"
    finally:
        release.set()
        assert transcribe._dependency_future.result(timeout=3) is dependencies


def test_concurrent_callers_join_one_import(monkeypatch, isolated_preparation):
    calls = []
    dependencies = (object(), object(), object())
    monkeypatch.setattr(transcribe, "_import_dependencies", lambda: calls.append(1) or dependencies)
    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = list(executor.map(lambda _: transcribe.start_dependency_preparation(), range(8)))
    assert all(future is futures[0] for future in futures)
    assert futures[0].result(timeout=3) is dependencies
    assert calls == [1]


def test_concurrent_model_callers_construct_only_one_backend(monkeypatch, isolated_preparation):
    entered = threading.Event()
    release = threading.Event()
    backend = object()
    loads = []

    def load():
        loads.append(1)
        entered.set()
        assert release.wait(3)
        return backend

    monkeypatch.setattr(transcribe, "_load_backend", load)
    with ThreadPoolExecutor(max_workers=2) as executor:
        first = executor.submit(transcribe.get_model)
        assert entered.wait(3)
        second = executor.submit(transcribe.get_model)
        release.set()
        assert first.result(timeout=3) is second.result(timeout=3) is backend
    assert loads == [1]


def test_import_failure_blocks_recording_with_original_reason(monkeypatch, isolated_preparation):
    def broken():
        raise ImportError("missing Cohere dependency")

    monkeypatch.setattr(transcribe, "_import_dependencies", broken)
    monkeypatch.setattr(transcribe, "_load_backend", transcribe._CohereBackend)
    future = transcribe.start_dependency_preparation()
    with pytest.raises(ImportError, match="missing Cohere dependency"):
        future.result(timeout=3)
    blocks, message = main._check_transcription_model()
    assert blocks
    assert "ImportError: missing Cohere dependency" in message
    assert transcribe._model is None


@pytest.mark.parametrize("flag", ["--list", "--list-sessions", "--prune-backups", "--dashboard"])
def test_explicit_non_recording_commands_skip_preparation(flag, monkeypatch):
    def forbidden():
        pytest.fail("a non-recording command prepared transcription dependencies")

    monkeypatch.setattr(transcribe, "start_dependency_preparation", forbidden)
    monkeypatch.setattr(main, "list_session_backups", lambda: None)
    monkeypatch.setattr(main, "prune_backups", lambda *args: None)
    import dashboard
    monkeypatch.setattr(dashboard, "serve", lambda: None)
    monkeypatch.setattr(sys, "argv", ["main.py", flag])
    main.run()


def test_another_menu_action_does_not_wait_for_unused_imports():
    # A blocked import would keep a ThreadPoolExecutor process alive at interpreter exit.
    script = """
import sys, threading
sys.path.insert(0, 'src')
import main
main.transcribe._import_dependencies = lambda: threading.Event().wait(60)
main.run_study_guide = lambda code: print('study guide selected', flush=True)
main.choose_notes_class = lambda args: 'TEST 100'
import builtins
builtins.input = lambda prompt: '2'
sys.argv = ['main.py']
main.run()
"""
    result = subprocess.run([sys.executable, "-B", "-c", script], cwd=ROOT,
                            capture_output=True, text=True, timeout=10)
    assert result.returncode == 0, result.stderr
    assert "study guide selected" in result.stdout


def test_importing_main_does_not_load_word_or_model_dependencies():
    script = """
import sys, json
sys.path.insert(0, 'src')
import main
main.transcribe.unload_model()
print(json.dumps([name for name in ('docx', 'torch', 'transformers') if name in sys.modules]))
"""
    result = subprocess.run([sys.executable, "-B", "-c", script], cwd=ROOT,
                            capture_output=True, text=True, timeout=10, check=True)
    assert json.loads(result.stdout) == []


@pytest.mark.parametrize("export", ["append_section", "save_study_guide", "rebuild"])
def test_lazy_word_exports_still_create_formatted_documents(export, tmp_path, monkeypatch):
    import docx_export
    from docx import Document

    monkeypatch.setattr(docx_export, "NOTES_DIR", tmp_path)
    path = getattr(docx_export, export)("TEST 100", "Test", "## Lecture\n- **Mitosis** makes cells.")
    document = Document(path)
    assert any(paragraph.text == "Lecture" for paragraph in document.paragraphs)
    bullet = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Mitosis makes cells.")
    assert bullet.runs[0].bold
    assert bullet.style.name == "List Bullet"
